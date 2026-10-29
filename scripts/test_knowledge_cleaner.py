"""验收上传文件的本地清洗。使用临时目录，不写入正式资料。"""

from __future__ import annotations

import ast
import hashlib
import io
import json
import os
import sqlite3
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
sys.path.insert(0, str(BACKEND))

TEMP = Path(tempfile.mkdtemp(prefix="knowledge-clean-"))
os.environ["KNOWLEDGE_FILES_DB"] = str(TEMP / "files.sqlite3")
os.environ["KNOWLEDGE_UPLOAD_DIR"] = str(TEMP / "uploads")
os.environ["KNOWLEDGE_CLEANED_DIR"] = str(TEMP / "cleaned")
os.environ["KNOWLEDGE_EXTRACTED_DIR"] = str(TEMP / "extracted")
os.environ["KNOWLEDGE_IMPORT_PREVIEW_DIR"] = str(TEMP / "previews")
os.environ["KNOWLEDGE_GRAPH_IMPORTS_DB"] = str(TEMP / "imports.sqlite3")
os.environ["STARTUP_TOOL_SELFTEST"] = "0"

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

CHECKS: list[bool] = []
EXCEL = ROOT / "data" / "official_dataset" / "钢铁设备数据集_开发集.xlsx"
GRAPH = ROOT / "data" / "processed" / "graph" / "graph.json"
SQLITE = ROOT / "data" / "runtime" / "monitoring_events.sqlite3"
OFFICIAL = ROOT / "data" / "official"
EXPECTED_XLSX = "1e9caeba1fe526f3ad1b5dd6445a9dbcc02fa22ec02273d9863ebc84e3e30f87"
DEFAULTS = (
    ROOT / "data" / "runtime" / "knowledge_files.sqlite3",
    ROOT / "data" / "runtime" / "knowledge_uploads",
    ROOT / "data" / "runtime" / "knowledge_cleaned",
    ROOT / "data" / "runtime" / "knowledge_extracted",
    ROOT / "data" / "runtime" / "knowledge_import_previews",
    ROOT / "data" / "runtime" / "knowledge_graph_imports.sqlite3",
)


def check(name: str, ok: bool, detail: str = "") -> None:
    CHECKS.append(bool(ok))
    print(("PASS" if ok else "FAIL"), name, detail)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def upload(client: TestClient, filename: str, content: bytes):
    return client.post(
        "/api/knowledge/files",
        params={"filename": filename},
        content=content,
        headers={"Content-Type": "application/octet-stream"},
    )


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("故障现象", level=1)
    document.add_paragraph("电机振动 4.2 mm/s，温度 75°C。")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "测点"
    table.rows[0].cells[1].text = "数值"
    table.rows[1].cells[0].text = "驱动端"
    table.rows[1].cells[1].text = "75"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def stored_json(file_id: str) -> dict:
    return json.loads((TEMP / "cleaned" / f"{file_id}.json").read_text(encoding="utf-8"))


def main() -> None:
    before = {
        "excel": file_hash(EXCEL),
        "graph": file_hash(GRAPH),
        "sqlite": file_hash(SQLITE) if SQLITE.is_file() else None,
        "official": {path.name: file_hash(path) for path in OFFICIAL.glob("*.docx") if not path.name.startswith("~$")},
        "defaults": [path.exists() for path in DEFAULTS],
    }
    modules = ast.parse((BACKEND / "app" / "services" / "knowledge_cleaner.py").read_text(encoding="utf-8"))
    imported: set[str] = set()
    for node in ast.walk(modules):
        if isinstance(node, ast.Import):
            imported.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.update(node.module.split("."))
    banned = imported & {"deepseek", "requests", "httpx", "monitoring_events", "anomaly_detection"}
    check("清洗模块没有调用 DeepSeek 或监测库", not banned, str(sorted(banned)))

    client = TestClient(app)
    text = "故障现象\n轴承温度达到 80°C。\n".encode()
    uploaded = upload(client, "轴承维修记录.txt", text)
    saved = uploaded.json()
    document = stored_json(saved["file_id"]) if uploaded.status_code == 200 else {}
    joined = "\n".join(item["text"] for item in document.get("sections", []))
    check("中文 TXT 自动生成 JSON 且不乱码", uploaded.status_code == 200 and saved["processing_status"] == "cleaned" and "轴承温度达到 80°C" in joined, saved.get("processing_status", ""))
    check("TXT 保留行号", any(item["locator"]["type"] == "line" and item["locator"]["index"] == 2 for item in document.get("sections", [])))

    word = upload(client, "轴承维修记录.docx", docx_bytes())
    word_body = word.json()
    word_doc = stored_json(word_body["file_id"]) if word.status_code == 200 else {}
    sections = word_doc.get("sections", [])
    headings = [item for item in sections if item["kind"] == "heading"]
    tables = [item for item in sections if item["kind"] == "table"]
    check(
        "DOCX 提取章节、正文和表格并保留位置",
        word.status_code == 200
        and headings
        and headings[0]["text"] == "故障现象"
        and headings[0]["locator"]["type"] == "paragraph"
        and any("4.2 mm/s" in item["text"] for item in sections)
        and tables
        and tables[0]["locator"]["type"] == "table"
        and tables[0]["table"]["rows"][1] == ["驱动端", "75"],
        word.text[:300],
    )

    csv_file = upload(client, "测点.csv", "测点,数值,单位\n驱动端,75,°C\n".encode())
    csv_doc = stored_json(csv_file.json()["file_id"])
    check("CSV 行内容一致", csv_doc["sections"][1]["cells"] == ["驱动端", "75", "°C"] and csv_doc["sections"][1]["locator"]["index"] == 2)

    md_file = upload(client, "记录.md", "# 故障现象\n振动 1.5 mm/s\n".encode())
    md_doc = stored_json(md_file.json()["file_id"])
    check("Markdown 标题和正文一致", md_doc["sections"][0]["heading"] == "故障现象" and "1.5 mm/s" in md_doc["sections"][1]["text"])

    json_file = upload(client, "记录.json", json.dumps({"设备": "棒线主轧机电机", "振动": 3.017}, ensure_ascii=False).encode())
    json_doc = stored_json(json_file.json()["file_id"])
    paths = {item["locator"]["path"]: item["text"] for item in json_doc["sections"]}
    check("JSON 字段路径和数值一致", paths.get("$.设备") == "棒线主轧机电机" and paths.get("$.振动") == "3.017")

    pdf = upload(client, "说明.pdf", b"%PDF-1.4\n")
    pdf_body = pdf.json()
    check(
        "PDF 明确失败且不生成清洗结果",
        pdf.status_code == 200
        and pdf_body["processing_status"] == "cleaning_failed"
        and "暂未完成" in pdf_body["error_message"]
        and not (TEMP / "cleaned" / f"{pdf_body['file_id']}.json").exists()
        and (TEMP / "uploads" / pdf_body["file_id"]).is_file(),
    )

    original = (TEMP / "uploads" / saved["file_id"]).read_bytes()
    first_json = (TEMP / "cleaned" / f"{saved['file_id']}.json").read_bytes()
    again = client.post(f"/api/knowledge/files/{saved['file_id']}/clean")
    second = client.post(f"/api/knowledge/files/{saved['file_id']}/clean")
    second_json = (TEMP / "cleaned" / f"{saved['file_id']}.json").read_bytes()
    check(
        "重复清洗不损坏 JSON 和原始文件",
        again.status_code == 200
        and second.status_code == 200
        and second.json()["processing_status"] == "cleaned"
        and json.loads(second_json)["sections"][1]["text"] == "轴承温度达到 80°C。"
        and (TEMP / "uploads" / saved["file_id"]).read_bytes() == original
        and b"file_id" in first_json
        and b"file_id" in second_json,
    )

    connection = sqlite3.connect(os.environ["KNOWLEDGE_FILES_DB"])
    status = connection.execute("SELECT processing_status FROM knowledge_files WHERE file_id = ?", (saved["file_id"],)).fetchone()
    connection.close()
    preview = client.get(f"/api/knowledge/files/{saved['file_id']}/cleaned")
    preview_body = preview.json()
    check("重新打开数据库后清洗状态仍在", status == ("cleaned",))
    check(
        "预览接口返回实际结果且不含磁盘路径",
        preview.status_code == 200
        and preview_body["preview"]["sections"][0]["text"] == "故障现象"
        and preview_body["in_graph"] is False
        and "knowledge_cleaned" not in preview.text
        and str(TEMP) not in preview.text,
    )
    long_text = "".join(f"第{index}行 振动 {index} mm/s\n" for index in range(1, 26)).encode()
    long_file = upload(client, "长记录.txt", long_text)
    long_id = long_file.json()["file_id"]
    long_preview = client.get(f"/api/knowledge/files/{long_id}/cleaned").json()
    long_stored = stored_json(long_id)
    check(
        "大文件接口只返回有限预览，完整结果仍保存",
        long_preview["preview_truncated"] is True
        and len(long_preview["preview"]["sections"]) == 20
        and len(long_stored["sections"]) == 25,
    )
    missing = client.get("/api/knowledge/files/abcdefabcdefabcdefabcdefabcdefab/cleaned")
    check("不存在的文件返回 404", missing.status_code == 404)
    pending = client.get(f"/api/knowledge/files/{pdf_body['file_id']}/cleaned")
    check("未清洗成功时不返回虚构 JSON", pending.status_code == 200 and pending.json()["preview"] is None)

    graph = json.loads(GRAPH.read_text(encoding="utf-8"))
    check("图谱仍是 252 个节点和 464 条关系", len(graph["nodes"]) == 252 and len(graph["edges"]) == 464 and file_hash(GRAPH) == before["graph"])
    official_after = {path.name: file_hash(path) for path in OFFICIAL.glob("*.docx") if not path.name.startswith("~$")}
    check("官方 Word 和 Excel 没有变化", official_after == before["official"] and file_hash(EXCEL) == EXPECTED_XLSX)
    check("监测事件库没有被写入", (file_hash(SQLITE) if SQLITE.is_file() else None) == before["sqlite"])
    check("测试没有创建正式清洗目录", [path.exists() for path in DEFAULTS] == before["defaults"])

    print(f"{sum(CHECKS)}/{len(CHECKS)} passed")
    if not all(CHECKS):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
