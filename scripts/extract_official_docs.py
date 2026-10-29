"""把 data/official 中的 Word 知识库解析为 Markdown。

只在本机读取 .docx，不访问网络，不调用 DeepSeek。
不修改、移动或删除原始 Word 文件。
"""

from __future__ import annotations

import hashlib
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "official"
OUTPUT_DIR = ROOT / "data" / "processed" / "markdown"
MANIFEST_PATH = ROOT / "data" / "processed" / "manifest.json"

HEADING_LEVELS = {
    "Title": 1,
    "Heading 1": 2,
    "Heading 2": 3,
    "Heading 3": 4,
    "Heading 4": 5,
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text.strip()


def style_name(paragraph: Paragraph) -> str:
    if paragraph.style is None or not paragraph.style.name:
        return "Normal"
    return paragraph.style.name


def cell_text(cell) -> str:
    parts = [paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()]
    text = "<br>".join(parts)
    return text.replace("|", "\\|")


def table_markdown(table: Table) -> str:
    rows = [[cell_text(cell) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def render_paragraph(paragraph: Paragraph, unknown_styles: set[str]) -> str | None:
    text = paragraph_text(paragraph)
    if not text:
        return None
    name = style_name(paragraph)
    level = HEADING_LEVELS.get(name)
    if level is not None:
        return f"{'#' * level} {text}"
    if name == "List Bullet":
        return f"- {text}"
    if name != "Normal":
        unknown_styles.add(name)
    return text


def unsupported_items(document: Document) -> list[str]:
    xml = document.element.body.xml
    found: list[str] = []
    if "a:blip" in xml or "w:drawing" in xml or "w:pict" in xml:
        found.append("文档包含图片，图片内容未提取")
    if "m:oMath" in xml:
        found.append("文档包含公式，公式内容未提取")
    if "w:txbxContent" in xml:
        found.append("文档包含文本框，文本框内容未按正文顺序提取")
    if "w:object" in xml or "w:oleObject" in xml:
        found.append("文档包含嵌入对象，对象内容未提取")
    return found


def expected_texts(document: Document) -> tuple[list[str], list[str]]:
    paragraphs: list[str] = []
    cells: list[str] = []
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = paragraph_text(block)
            if text:
                paragraphs.append(text)
        else:
            for row in block.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            cells.append(text)
    return paragraphs, cells


def text_present(markdown: str, text: str) -> bool:
    return text in markdown or text.replace("|", "\\|") in markdown


def source_preamble(source_name: str, source_path: str) -> str:
    return "\n".join(
        [
            "> 原始文档名称：" + source_name,
            "> 来源位置：" + source_path,
            "",
        ]
    )


def convert_document(source: Path) -> dict:
    source_path = source.relative_to(ROOT).as_posix()
    output = OUTPUT_DIR / f"{source.stem}.md"
    output_path = output.relative_to(ROOT).as_posix()
    before_hash = file_sha256(source)
    record = {
        "source_name": source.name,
        "source_path": source_path,
        "output_path": output_path,
        "section_count": 0,
        "title_count": 0,
        "heading_1_count": 0,
        "heading_2_count": 0,
        "table_count": 0,
        "paragraph_count": 0,
        "list_item_count": 0,
        "status": "failed",
        "limitations": [],
        "notes": [],
        "exceptions": [],
    }
    try:
        document = Document(str(source))
        unknown_styles: set[str] = set()
        blocks: list[str] = [source_preamble(source.name, source_path)]
        for block in iter_blocks(document):
            if isinstance(block, Paragraph):
                name = style_name(block)
                rendered = render_paragraph(block, unknown_styles)
                if rendered is None:
                    continue
                record["paragraph_count"] += 1
                if name == "Title":
                    record["title_count"] += 1
                elif name == "Heading 1":
                    record["heading_1_count"] += 1
                    record["section_count"] += 1
                elif name == "Heading 2":
                    record["heading_2_count"] += 1
                    record["section_count"] += 1
                elif name == "List Bullet":
                    record["list_item_count"] += 1
                blocks.append(rendered)
                blocks.append("")
            else:
                rendered_table = table_markdown(block)
                record["table_count"] += 1
                if rendered_table:
                    blocks.append(rendered_table)
                    blocks.append("")
        limitations = unsupported_items(document)
        if unknown_styles:
            names = "、".join(sorted(unknown_styles))
            limitations.append(f"存在未单独映射的段落样式：{names}。这些段落已按原文保留")
        if record["list_item_count"]:
            record["notes"].append("List Bullet 样式的项目符号在原文字符中不存在，Markdown 用列表符号表示，文字未改写")
        markdown = "\n".join(blocks).rstrip() + "\n"
        output.write_text(markdown, encoding="utf-8")
        paragraphs, cells = expected_texts(document)
        missing_paragraphs = sum(1 for text in paragraphs if not text_present(markdown, text))
        missing_cells = sum(1 for text in cells if not text_present(markdown, text))
        separator_count = sum(1 for line in markdown.splitlines() if line.startswith("| ---"))
        if missing_paragraphs:
            limitations.append(f"有 {missing_paragraphs} 段正文未写入 Markdown")
        if missing_cells:
            limitations.append(f"有 {missing_cells} 个表格单元格未写入 Markdown")
        if separator_count != record["table_count"]:
            limitations.append("Markdown 表格数量与 Word 表格数量不一致")
        record["limitations"] = limitations
        hard_failure = missing_paragraphs or missing_cells or separator_count != record["table_count"]
        unparsed = any(
            item.startswith("文档包含")
            for item in limitations
        )
        if hard_failure:
            record["status"] = "failed"
        elif unparsed:
            record["status"] = "partial"
        else:
            record["status"] = "success"
    except Exception as exc:
        record["exceptions"].append(type(exc).__name__)
        record["status"] = "failed"
    finally:
        if file_sha256(source) != before_hash:
            record["exceptions"].append("原始 Word 文件哈希发生变化")
            record["status"] = "failed"
    return record


def main() -> None:
    sources = sorted(SOURCE_DIR.glob("*.docx"))
    if len(sources) != 6:
        raise SystemExit(f"期望读取 6 份 Word，实际找到 {len(sources)} 份")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    documents = [convert_document(source) for source in sources]
    manifest = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_directory": SOURCE_DIR.relative_to(ROOT).as_posix(),
        "output_directory": OUTPUT_DIR.relative_to(ROOT).as_posix(),
        "document_count": len(documents),
        "table_count": sum(item["table_count"] for item in documents),
        "documents": documents,
    }
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    failed = [item["source_name"] for item in documents if item["status"] == "failed"]
    print(f"documents={len(documents)} tables={manifest['table_count']} failed={len(failed)}")
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
